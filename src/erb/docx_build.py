"""`erb build TICKER`: assemble the initiating coverage report (.docx).

Inputs in coverage/<TICKER>/:
  sections/00_cover.md ... 10_appendix.md   narrative (agent-written Markdown, see `erb scaffold`)
  model.json, charts/*.png, facts/company.json

Markdown subset: `#`/`##`/`###` headings, paragraphs, `- ` bullets, `**bold**`,
`*italic*`, pipe tables, `![Title](file.png)` images (from charts/) with an optional
`Source: ...` line after, `{{token}}` exhibits on their own line (see EXHIBITS),
and HTML comments (dropped). `[S12]` source tags are stripped from the prose
(they live in sources.md). `[VERIFY: ...]` is kept and highlighted so no one can miss it.
"""

from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import fmt
from .config import ASSETS_DIR, GUIDES_DIR, coverage_dir

BRAND = yaml.safe_load((ASSETS_DIR / "brand.yaml").read_text())
COL = BRAND["colors"]
HEAD_FONT, BODY_FONT = BRAND["fonts"]["heading"], BRAND["fonts"]["body"]
AUTHOR = {"name": "Ethan Stott", "email": "stott@consciousinvestments.org"}
PAGE_W, MARGIN = 8.5, 0.6
TEXT_W = PAGE_W - 2 * MARGIN


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


# ---------------------------------------------------------------- low-level XML helpers

def shade(cell, hex_: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_.lstrip("#"))
    tcPr.append(shd)


def table_borders(table, horizontal: str | None = COL["grid"], outer: str | None = None, size: int = 4) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        colour = horizontal if edge in ("insideH", "bottom", "top") else None
        if edge in ("top", "bottom") and outer:
            colour = outer
        if colour:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:color"), colour.lstrip("#"))
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def cell_margins(table, top=30, bottom=30, left=70, right=70) -> None:
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def widths(table, inches: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, inches):
            cell.width = Inches(w)


def para_border(p, edge: str = "bottom", colour: str = COL["ink"], size: int = 6, space: int = 2) -> None:
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), colour.lstrip("#"))
    pbdr.append(el)
    pPr.append(pbdr)


def add_field(p, instr: str, placeholder: str = "") -> None:
    def fld(kind):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        return r
    fld("begin")
    r = p.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {instr} "
    r._r.append(t)
    fld("separate")
    p.add_run(placeholder)
    fld("end")


def keep_with_next(p) -> None:
    p.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- styles

def setup_styles(doc: Document) -> None:
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(COL["ink"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.space_after = Pt(4)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.07

    for name, size, before, after in (("Heading 1", 12.5, 10, 4), ("Heading 2", 10, 8, 2), ("Heading 3", 9.5, 6, 1)):
        h = st[name]
        h.font.name = HEAD_FONT
        h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
        h.element.rPr.rFonts.set(qn("w:ascii"), HEAD_FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), HEAD_FONT)
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = rgb(COL["ink"])
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    st["Heading 1"].font.all_caps = True

    from docx.enum.style import WD_STYLE_TYPE
    for name, bold, caps, indent in (("toc 1", True, True, 0.0), ("toc 2", False, False, 0.2)):
        try:
            ts = st[name]
        except KeyError:
            ts = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            ts.base_style = st["Normal"]
        ts.style_id = name.upper().replace(" ", "")  # Word's built-in ids are TOC1 / TOC2
        ts.font.bold, ts.font.all_caps, ts.font.size = bold, caps, Pt(9)
        ts.paragraph_format.left_indent = Inches(indent)
        ts.paragraph_format.space_after = Pt(2 if bold else 1)
        ts.paragraph_format.space_before = Pt(4 if bold else 0)

    for name in ("List Bullet",):
        s = st[name]
        s.font.name = BODY_FONT
        s.font.size = Pt(9.5)
        s.paragraph_format.space_after = Pt(3)


# ---------------------------------------------------------------- inline markdown

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\[VERIFY:[^\]]*\])")
SOURCE_TAG = re.compile(r"\s?\[(?:S\d+(?:\s*,\s*S\d+)*|M)\]")


def add_inline(p, text: str, size: float | None = None, color: str | None = None, bold: bool = False) -> None:
    text = SOURCE_TAG.sub("", text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("[VERIFY:"):
            r = p.add_run(part)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            r = p.add_run(part)
            r.bold = bold or None
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)


def body_para(container, text: str, justify: bool = True, size: float | None = None):
    p = container.add_paragraph()
    add_inline(p, text, size=size)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def small(container, text: str, color: str = COL["text_secondary"], size: float = 7, italic: bool = False,
          align=None, bold: bool = False):
    p = container.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = rgb(color)
    r.italic = italic
    r.bold = bold
    p.paragraph_format.space_after = Pt(1)
    if align:
        p.alignment = align
    return p


# ---------------------------------------------------------------- tables

def data_table(container, header: list[str], rows: list[list[str]], col_w: list[float] | None = None,
               bold_rows: set[int] = frozenset(), highlight: set[tuple[int, int]] = frozenset(),
               size: float = 7.5, first_col_left: bool = True, shade_rows: set[int] = frozenset()):
    t = container.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t, horizontal=COL["grid"], outer=COL["ink"])
    cell_margins(t, 20, 20, 60, 60)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        shade(c, COL["paper"])
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(size)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and first_col_left) else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = t.rows[i].cells[j]
            p = c.paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(size)
            r.bold = (i - 1) in bold_rows or None
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 and first_col_left) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            if (i - 1, j) in highlight:
                shade(c, COL["paper"])
                r.bold = True
            elif (i - 1) in shade_rows:
                shade(c, COL["paper_light"])
    for row in t.rows:
        for c in row.cells:
            mark_size(c.paragraphs[0], size)
    if col_w:
        widths(t, col_w)
    return t


def mark_size(p, size: float) -> None:
    """Set the paragraph-mark font size so empty cells don't grow the row."""
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        pPr.append(rPr)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(int(size * 2)))
        rPr.append(el)


def spacer(doc, pts: float = 4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    mark_size(p, pts)
    return p


def exhibit_title(container, text: str):
    p = container.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.name = BODY_FONT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    keep_with_next(p)
    return p


def source_line(container, text: str):
    p = small(container, text if text.lower().startswith("source") else f"Source: {text}", size=6.5, italic=True)
    p.paragraph_format.space_after = Pt(6)
    return p


def image(container, path: str | Path, width: float = TEXT_W, source: str | None = None):
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = bool(source)
    if source:
        source_line(container, source)


def md_table(container, lines: list[str]):
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not rows:
        return
    data_table(container, [SOURCE_TAG.sub("", c).replace("**", "") for c in rows[0]],
               [[SOURCE_TAG.sub("", c).replace("**", "") for c in r] for r in rows[1:]])
    container.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- report context

class Ctx:
    def __init__(self, ticker: str):
        self.dir = coverage_dir(ticker)
        self.model = json.loads((self.dir / "model.json").read_text())
        self.company = json.loads((self.dir / "facts" / "company.json").read_text())
        idx = self.dir / "charts" / "index.json"
        self.charts = json.loads(idx.read_text()) if idx.exists() else {}
        self.ticker = self.model["ticker"]
        # EDGAR names are often ALL CAPS; prefer Yahoo's proper-case long name
        self.name = self.company["market"].get("name") or self.company["name"]
        self.cur = self.model["cover"]["currency"] or "USD"
        self.fin_cur = self.company["financial_currency"]
        self.front: dict = {}
        self.missing_tokens: list[str] = []

    def p(self, x) -> str:
        return fmt.price(x, self.cur)

    def m(self, x) -> str:
        return fmt.money(x, self.fin_cur)


# ---------------------------------------------------------------- exhibits ({{token}})

def ex_chart(key: str, source: str, width: float = TEXT_W):
    def render(doc, ctx: Ctx):
        path = ctx.charts.get(key)
        if not path:
            ctx.missing_tokens.append(key)
            return
        image(doc, path, width=width, source=source)
    return render


def ex_company_snapshot(doc, ctx: Ctx):
    f, m = ctx.front, ctx.company["market"]
    rows = [["Headquarters", f.get("hq") or m.get("hq") or "[VERIFY: HQ]",
             "Chief Executive Officer", f.get("ceo") or "[VERIFY: CEO]"],
            ["Number of Employees", f.get("employees") or fmt.number(m.get("employees"), 1),
             "GICS Sector", f.get("gics_sector") or m.get("sector") or ""],
            ["GICS Sub-Industry", f.get("gics_sub_industry") or m.get("industry") or "", "Fiscal Year End",
             ctx.company["edgar"].get("fiscal_year_end_mmdd", "")]]
    t = doc.add_table(rows=3, cols=4)
    table_borders(t, horizontal=COL["grid"])
    cell_margins(t)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            p = t.rows[i].cells[j].paragraphs[0]
            add_inline(p, str(v), size=7.5, bold=(j % 2 == 0))
            p.paragraph_format.space_after = Pt(0)
            if j % 2 == 0:
                shade(t.rows[i].cells[j], COL["paper_light"])
    widths(t, [1.35, 2.3, 1.5, 2.15])
    source_line(doc, "SEC filings, Yahoo Finance")


def ex_dcf(doc, ctx: Ctx):
    sc = ctx.model["scenarios"]
    order = ("bear", "base", "bull")
    fc = ctx.fin_cur
    a = ctx.model["assumptions"]
    term = a["terminal"]
    rows = [
        ["WACC"] + [fmt.pct(sc[s]["wacc"]["wacc"]) for s in order],
        ["Cost of equity"] + [fmt.pct(sc[s]["wacc"]["cost_of_equity"]) for s in order],
        ["Terminal method"] + [("Exit " + fmt.multiple(sc[s]["dcf"]["exit_multiple"]) + " EV/EBITDA")
                               if term.get("method", "exit_multiple") == "exit_multiple"
                               else f"Perpetuity g {fmt.pct(term.get('perpetual_growth'))}" for s in order],
        ["PV of cash flows"] + [fmt.money(sc[s]["dcf"]["pv_cash_flows"], fc) for s in order],
        ["PV of terminal value"] + [fmt.money(sc[s]["dcf"]["pv_terminal"], fc) for s in order],
        ["Enterprise value"] + [fmt.money(sc[s]["dcf"]["enterprise_value"], fc) for s in order],
        ["Less: net debt (negative = net cash)"] + [fmt.money(a["net_debt"], fc)] * 3,
        ["Equity value"] + [fmt.money(sc[s]["dcf"]["equity_value"], fc) for s in order],
        ["Implied perpetual growth"] + [fmt.pct(sc[s]["dcf"]["implied_perpetual_growth"]) for s in order],
        ["DCF price target"] + [ctx.p(sc[s]["methods"]["dcf"]) for s in order],
    ]
    for key, label in (("pe", "Forward P/E price target"), ("ev_ebitda", "EV/EBITDA price target"),
                       ("metric", (a.get("metric_multiple") or {}).get("name", "Metric") + " price target")):
        if key in sc["base"]["methods"]:
            rows.append([label] + [ctx.p(sc[s]["methods"][key]) for s in order])
    w = sc["base"]["weights"]
    rows.append(["Blended price target (" + ", ".join(f"{k.upper().replace('_', '/')} {v:.0%}" for k, v in w.items()) + ")"]
                + [ctx.p(sc[s]["price_target"]) for s in order])
    rows.append(["Implied return"] + [fmt.pct(sc[s]["price_return"]) for s in order])
    exhibit_title(doc, f"{ctx.ticker} Valuation Summary")
    data_table(doc, ["", "Bear Case", "Base Case", "Bull Case"], rows, [3.1, 1.4, 1.4, 1.4],
               bold_rows={len(rows) - 2, len(rows) - 1}, highlight={(len(rows) - 2, 2)})
    source_line(doc, "Conscious Investments model")


def ex_sensitivity(doc, ctx: Ctx):
    s = ctx.model["sensitivity"]
    method_exit = s["col_label"].startswith("Exit")
    cols = [fmt.multiple(c) if method_exit else fmt.pct(c) for c in s["cols"]]
    rows = [[fmt.pct(w)] + [ctx.p(v) for v in row] for w, row in zip(s["rows_wacc"], s["values"])]
    mid_r, mid_c = len(rows) // 2, len(cols) // 2 + 1
    exhibit_title(doc, f"{ctx.ticker} DCF Price Target Sensitivity: WACC vs. {s['col_label']}")
    data_table(doc, ["WACC \\ " + s["col_label"]] + cols, rows, [1.6] + [1.1] * len(cols), highlight={(mid_r, mid_c)})
    source_line(doc, "Conscious Investments model")


def ex_assumptions(doc, ctx: Ctx):
    a = ctx.model["assumptions"]
    proj = ctx.model["projections"]["base"]
    years = [fmt.fy(int(r["fy"])) + "E" for r in proj]
    rows = [["Revenue growth"] + [fmt.pct(r["revenue_growth"]) for r in proj],
            ["EBITDA margin"] + [fmt.pct(r["ebitda_margin"]) for r in proj],
            ["D&A % of revenue"] + [fmt.pct(r["d_and_a"] / r["revenue"]) for r in proj],
            ["CapEx % of revenue"] + [fmt.pct(r["capex"] / r["revenue"]) for r in proj],
            ["Tax rate"] + [fmt.pct(r["tax_rate"]) for r in proj]]
    exhibit_title(doc, f"{ctx.ticker} Base Case Model Assumptions")
    data_table(doc, [""] + years, rows, [1.8] + [1.1] * len(years))
    spacer(doc)
    w = a["wacc"]
    rows2 = [["Risk-free rate", fmt.pct(w["risk_free"]), "Beta", f"{w['beta']:.2f}"],
             ["Equity risk premium", fmt.pct(w["erp"]), "Pre-tax cost of debt", fmt.pct(w["cost_of_debt_pretax"])],
             ["Target forward P/E", fmt.multiple(a["multiples"].get("forward_pe")),
              "Target EV/EBITDA", fmt.multiple(a["multiples"].get("ev_ebitda"))],
             ["Exit EV/EBITDA", fmt.multiple(a["terminal"].get("exit_ev_ebitda")), "Horizon",
              f"{a.get('horizon_months', 15)} months"]]
    data_table(doc, ["Input", "Value", "Input", "Value"], rows2, [1.9, 1.75, 1.9, 1.75])
    source_line(doc, "Conscious Investments model; risk-free rate from 10Y UST; beta from Yahoo Finance")


def ex_financial_summary(doc, ctx: Ctx):
    annual_path = ctx.dir / "facts" / "financials_annual.csv"
    import pandas as pd
    a = pd.read_csv(annual_path, index_col="period").dropna(subset=["revenue"]).tail(3)
    proj = ctx.model["projections"]["base"]
    fc = ctx.fin_cur
    hdr = ["Summary"] + [str(int(y)) for y in a["fy"]] + [f"{int(r['fy'])}e" for r in proj]
    seg_keys = [k for k in proj[0] if k.startswith("seg:")]

    def rowvals(hist_col, proj_col, f):
        return [f(v) for v in a[hist_col]] + [f(r[proj_col]) for r in proj] if hist_col in a else \
            ["N/A"] * len(a) + [f(r[proj_col]) for r in proj]

    rev_all = list(a["revenue"]) + [r["revenue"] for r in proj]
    growth = ["N/A"] + [fmt.pct(rev_all[i] / rev_all[i - 1] - 1) for i in range(1, len(rev_all))]
    eps_all = list(a["eps_diluted"]) + [r["eps"] for r in proj]
    rows = [["Total revenue"] + [fmt.money(v, fc) for v in rev_all], ["  % growth YoY"] + growth]
    seg_hist = {}
    seg_path = ctx.dir / "facts" / "segments.csv"
    if seg_keys and seg_path.exists() and seg_path.stat().st_size > 1:
        sg = pd.read_csv(seg_path)
        sg = sg[(sg["dimension"] == "segment") & (sg["concept"] == "revenue") & sg["days"].between(350, 380)]
        ends = {str(e)[:10]: fy for e, fy in zip(a["end"], a["fy"])}
        for r in sg.itertuples():
            if str(r.end)[:10] in ends:
                seg_hist[(r.label, ends[str(r.end)[:10]])] = r.value
    for k in seg_keys:
        hist = [fmt.money(seg_hist[(k[4:], fy)], fc) if (k[4:], fy) in seg_hist else "" for fy in a["fy"]]
        rows.append([f"  {k[4:]}"] + hist + [fmt.money(r[k], fc) for r in proj])
    rows += [["EBITDA"] + rowvals("ebitda", "ebitda", lambda v: fmt.money(v, fc)),
             ["  EBITDA margin"] + rowvals("ebitda_margin", "ebitda_margin", fmt.pct),
             ["Net income"] + rowvals("net_income", "net_income", lambda v: fmt.money(v, fc)),
             ["Diluted EPS"] + [fmt.price(v, fc) for v in eps_all],
             ["  % growth YoY"] + ["N/A"] + [fmt.pct(eps_all[i] / eps_all[i - 1] - 1) if eps_all[i - 1] > 0 else "N/A"
                                             for i in range(1, len(eps_all))],
             ["Free cash flow*"] + rowvals("fcf", "ufcf", lambda v: fmt.money(v, fc))]
    exhibit_title(doc, f"{ctx.ticker} Financial Summary ({fc})")
    n = len(hdr) - 1
    data_table(doc, hdr, rows, [1.5] + [(TEXT_W - 1.5) / n] * n, size=7,
               shade_rows={i for i, r in enumerate(rows) if not r[0].startswith("  ")})
    source_line(doc, "SEC filings (actuals), Conscious Investments model (estimates). "
                     "*Actual FCF = CFO less CapEx; estimates are unlevered FCF. Estimates are base case.")


def ex_comps(doc, ctx: Ctx):
    comps = ctx.model.get("comps") or []
    if not comps:
        ctx.missing_tokens.append("comps (set peers: in assumptions.yaml and rerun erb model)")
        return
    stats = ctx.model["comp_stats"]

    def f_m(v):
        return fmt.money(v) if v is not None else "N/A"

    cols = [("Mkt Cap", "market_cap", f_m), ("EV", "enterprise_value", f_m), ("Rev Growth", "revenue_growth", fmt.pct),
            ("EBITDA Mgn", "ebitda_margin", fmt.pct), ("EV/Sales", "ev_sales", fmt.multiple),
            ("EV/EBITDA", "ev_ebitda", fmt.multiple), ("P/E LTM", "pe_ltm", fmt.multiple), ("P/E NTM", "pe_ntm", fmt.multiple)]
    rows = [[c.get("name") or c["ticker"], c["ticker"]] + [f(c.get(k)) for _, k, f in cols] for c in comps]
    for label, key in (("High", "max"), ("Mean", "mean"), ("Median", "median"), ("Low", "min")):
        rows.append([label, ""] + [f(stats[key].get(k)) for _, k, f in cols])
    exhibit_title(doc, f"{ctx.ticker} Peer Valuation Comparison")
    n = len(comps)
    data_table(doc, ["Company", "Ticker"] + [c[0] for c in cols], rows, [1.75, 0.55] + [0.625] * 8, size=6.8,
               bold_rows={0} | set(range(n, n + 4)), shade_rows={0})
    cols2 = [("ROE", "roe", fmt.pct), ("ROA", "roa", fmt.pct), ("Gross Mgn", "gross_margin", fmt.pct),
             ("Net Mgn", "net_margin", fmt.pct), ("Debt/Equity", "debt_to_equity", lambda v: fmt.multiple(v) if v else "N/A"),
             ("Beta", "beta", lambda v: f"{v:.2f}" if v is not None else "N/A"),
             ("Div Yield", "dividend_yield", fmt.pct), ("Current Ratio", "current_ratio", lambda v: f"{v:.2f}" if v else "N/A")]
    rows2 = [[c.get("name") or c["ticker"], c["ticker"]] + [f(c.get(k)) for _, k, f in cols2] for c in comps]
    for label, key in (("High", "max"), ("Mean", "mean"), ("Median", "median"), ("Low", "min")):
        rows2.append([label, ""] + [f(stats[key].get(k)) for _, k, f in cols2])
    exhibit_title(doc, f"{ctx.ticker} Peer Returns, Margins and Balance Sheet")
    data_table(doc, ["Company", "Ticker"] + [c[0] for c in cols2], rows2, [1.75, 0.55] + [0.625] * 8, size=6.8,
               bold_rows={0} | set(range(n, n + 4)), shade_rows={0})
    source_line(doc, "Yahoo Finance; statistics exclude the subject company and negative multiples")


EXHIBITS = {
    "price_targets": ex_chart("price_targets", "Conscious Investments model, Yahoo Finance"),
    "band_pe": ex_chart("band_ltm_pe", "SEC filings, Yahoo Finance; LTM P/E is point-in-time (as filed)"),
    "band_ev_ebitda": ex_chart("band_ltm_ev_ebitda", "SEC filings, Yahoo Finance; LTM EV/EBITDA is point-in-time"),
    "segment_mix": ex_chart("mix_segment", "SEC filings (10-K segment note)", width=3.6),
    "geography_mix": ex_chart("mix_geography", "SEC filings", width=3.6),
    "product_mix": ex_chart("mix_product", "SEC filings", width=3.6),
    "financial_grid": ex_chart("financial_grid", "SEC filings, Conscious Investments model"),
    "football_field": ex_chart("football_field", "Conscious Investments model, Yahoo Finance"),
    "company_snapshot": ex_company_snapshot,
    "dcf": ex_dcf,
    "sensitivity": ex_sensitivity,
    "assumptions": ex_assumptions,
    "financial_summary": ex_financial_summary,
    "comps": ex_comps,
}


# ---------------------------------------------------------------- markdown sections

def split_front(text: str) -> tuple[dict, str]:
    if text.startswith("---"):
        _, fm, body = text.split("---", 2)
        return yaml.safe_load(fm) or {}, body
    return {}, text


def render_md(doc, text: str, ctx: Ctx, scenario_headings: bool = False) -> None:
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    lines = text.splitlines()
    i, buf = 0, []

    def flush():
        if buf:
            body_para(doc, " ".join(s.strip() for s in buf))
            buf.clear()

    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            flush()
        elif re.fullmatch(r"\{\{\s*[\w]+\s*\}\}", s):
            flush()
            key = s.strip("{} ")
            fn = EXHIBITS.get(key)
            if fn:
                fn(doc, ctx)
            else:
                ctx.missing_tokens.append(key)
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            title = SOURCE_TAG.sub("", s.lstrip("#").strip())
            if scenario_headings and level == 2 and title.lower().split()[0] in ("bull", "base", "bear"):
                scenario_heading(doc, ctx, title.lower().split()[0])
            else:
                if level == 1:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE) if ctx_has_content(doc) else None
                doc.add_heading(title.upper() if level == 1 else title, level=min(level, 3))
        elif s.startswith("- ") or s.startswith("* "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, s[2:])
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif s.startswith("|"):
            flush()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            md_table(doc, block)
            continue
        elif m := re.fullmatch(r"!\[(.*?)\]\((.*?)\)", s):
            flush()
            title, fname = m.groups()
            path = (ctx.dir / "charts" / fname) if not Path(fname).is_absolute() else Path(fname)
            src = None
            if i + 1 < len(lines) and lines[i + 1].strip().lower().startswith("source"):
                src = lines[i + 1].strip()
                i += 1
            if title:
                exhibit_title(doc, title)
            if path.exists():
                image(doc, path, source=src)
            else:
                body_para(doc, f"[VERIFY: missing image {fname}]")
        else:
            buf.append(s)
        i += 1
    flush()


def ctx_has_content(doc) -> bool:
    return any(p.text.strip() for p in doc.paragraphs[-3:])


def scenario_heading(doc, ctx: Ctx, s: str) -> None:
    sc = ctx.model["scenarios"][s]
    p = doc.add_paragraph()
    r = p.add_run(f"{s.title()} Case Price Target: {ctx.p(sc['price_target'])}")
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = rgb(BRAND["colors"]["scenario"][s] if s != "base" else COL["ink"])
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    keep_with_next(p)
    p2 = doc.add_paragraph()
    r2 = p2.add_run(f"12–18 Month Target Return: {fmt.pct(sc['price_return'])}")
    r2.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = rgb(COL["text_secondary"])
    p2.paragraph_format.space_after = Pt(2)
    keep_with_next(p2)


# ---------------------------------------------------------------- page furniture

def header_footer(section, ctx: Ctx, report_date: str) -> None:
    section.different_first_page_header_footer = True
    hdr = section.header
    t = hdr.add_table(rows=1, cols=2, width=Inches(TEXT_W))
    table_borders(t, horizontal=None)
    widths(t, [TEXT_W / 2, TEXT_W / 2])
    left, right = t.rows[0].cells
    lp = left.paragraphs[0]
    lp.add_run().add_picture(str(ASSETS_DIR / BRAND["logo"]), height=Inches(0.42))
    r = lp.add_run("   CONSCIOUS INVESTMENTS")
    r.font.name = HEAD_FONT
    r.font.size = Pt(9)
    r.font.color.rgb = rgb(COL["ink"])
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for text, size, bold in ((report_date, 8, True), (f"{ctx.name} ({ctx.ticker})", 7.5, False),
                             ("Initiating Coverage Report", 7.5, False)):
        p = right.paragraphs[0] if not right.paragraphs[0].text else right.add_paragraph()
        rr = p.add_run(text)
        rr.font.size = Pt(size)
        rr.bold = bold
        rr.font.color.rgb = rgb(COL["text_secondary"] if not bold else COL["ink"])
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
    rule = hdr.add_paragraph()
    para_border(rule, "bottom", COL["ink"], 6)
    rule.paragraph_format.space_after = Pt(0)

    ftr = section.footer
    fp = ftr.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Conscious Investments  ·  See important disclosures at the end of this report  ·  Page ")
    run.font.size = Pt(7)
    run.font.color.rgb = rgb(COL["text_muted"])
    add_field(fp, "PAGE", "1")
    for r in fp.runs[1:]:
        r.font.size = Pt(7)
        r.font.color.rgb = rgb(COL["text_muted"])

    # cover-page footer: source note only
    cf = section.first_page_footer.paragraphs[0]
    rr = cf.add_run(f"Source: SEC filings, Yahoo Finance, Conscious Investments model. Prices as of the "
                    f"{fmt_date(ctx.model['assumptions']['as_of'])} close. Please see important disclosures "
                    f"at the end of this report.")
    rr.font.size = Pt(6.5)
    rr.italic = True
    rr.font.color.rgb = rgb(COL["text_muted"])


def fmt_date(iso: str) -> str:
    d = date.fromisoformat(str(iso)[:10])
    return f"{d.day} {d.strftime('%B %Y')}"


# ---------------------------------------------------------------- cover

def cover(doc, ctx: Ctx, front: dict, body: str, report_date: str) -> None:
    m, mo = ctx.model, ctx.model
    top = doc.add_table(rows=1, cols=2)
    table_borders(top, horizontal=None)
    widths(top, [1.3, TEXT_W - 1.3])
    lc, rc = top.rows[0].cells
    lc.paragraphs[0].add_run().add_picture(str(ASSETS_DIR / BRAND["logo_cover"]), width=Inches(1.05))
    rc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for text, size, font, bold, colour in (
            (report_date, 9, BODY_FONT, True, COL["ink"]),
            ("CONSCIOUS INVESTMENTS", 17, HEAD_FONT, False, COL["ink"]),
            ("Initiating Coverage Report", 10, HEAD_FONT, False, COL["text_secondary"])):
        p = rc.paragraphs[0] if not rc.paragraphs[0].text else rc.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.font.name, r.bold = Pt(size), font, bold
        r.font.color.rgb = rgb(colour)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(1)
    rule = doc.add_paragraph()
    para_border(rule, "bottom", COL["ink"], 12)
    rule.paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph()
    r = title.add_run(f"{ctx.name} ({ctx.ticker})")
    r.font.name, r.font.size, r.bold = HEAD_FONT, Pt(19), True
    title.paragraph_format.space_after = Pt(0)
    tag = doc.add_paragraph()
    r = tag.add_run(front.get("tagline") or "[VERIFY: tagline]")
    r.font.name, r.font.size, r.italic = HEAD_FONT, Pt(12), True
    r.font.color.rgb = rgb(COL["text_secondary"])
    tag.paragraph_format.space_after = Pt(6)

    grid = doc.add_table(rows=1, cols=2)
    table_borders(grid, horizontal=None)
    main_w, side_w = 4.75, TEXT_W - 4.75 - 0.05
    widths(grid, [main_w, side_w])
    left, side = grid.rows[0].cells
    left.paragraphs[0].paragraph_format.space_after = Pt(0)
    render_cover_body(left, body, ctx)
    sidebar(side, ctx, side_w)


def render_cover_body(cell, body: str, ctx: Ctx) -> None:
    body = re.sub(r"<!--.*?-->", "", body, flags=re.S)
    first = True
    for block in re.split(r"\n\s*\n", body.strip()):
        block = block.strip()
        if not block:
            continue
        if block.startswith("#"):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            r = p.add_run(block.lstrip("#").strip().upper())
            r.font.name, r.bold, r.font.size = HEAD_FONT, True, Pt(10)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            para_border(p, "bottom", COL["grid"], 4)
        else:
            for item in re.split(r"\n(?=[-*] )", block):
                is_bullet = item.startswith(("- ", "* "))
                text = " ".join(s.strip() for s in (item[2:] if is_bullet else item).splitlines())
                if first and not cell.paragraphs[0].text:
                    p = cell.paragraphs[0]
                    if is_bullet:
                        p.style = cell.part.document.styles["List Bullet"]
                else:
                    p = cell.add_paragraph(style="List Bullet" if is_bullet else None)
                add_inline(p, text, size=9)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(3)
                first = False
        first = False


def sidebar(cell, ctx: Ctx, w: float) -> None:
    shade(cell, COL["paper_light"])
    m, c = ctx.model, ctx.model["cover"]
    sc, rt = m["scenarios"], m["rating"]

    def label(text, first=False):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        r = p.add_run(text)
        r.font.size, r.bold = Pt(7), True
        r.font.color.rgb = rgb(COL["text_secondary"])
        p.paragraph_format.space_before = Pt(4 if not first else 0)
        p.paragraph_format.space_after = Pt(1)
        return p

    label("RATING", first=True)
    p = cell.add_paragraph()
    r = p.add_run(rt["rating"].upper())
    r.font.name, r.font.size, r.bold = HEAD_FONT, Pt(15), True
    p.paragraph_format.space_after = Pt(0)
    small(cell, f"vs. S&P 500 expected {fmt.pct(rt['benchmark_return'])} over {rt['horizon_months']} months", size=6.5)

    t = cell.add_table(rows=3, cols=4)
    table_borders(t, horizontal=COL["grid"], outer=COL["ink"])
    cell_margins(t, 15, 15, 20, 20)
    heads = ["Bear", "Current", "Base PT", "Bull"]
    vals = [sc["bear"]["price_target"], m["price"], sc["base"]["price_target"], sc["bull"]["price_target"]]
    rets = [fmt.pct(sc["bear"]["price_return"]), "", fmt.pct(sc["base"]["price_return"]), fmt.pct(sc["bull"]["price_return"])]
    for j in range(4):
        for i, text in enumerate((heads[j], ctx.p(vals[j]).replace(".00", ""), rets[j])):
            cc = t.rows[i].cells[j]
            pp = cc.paragraphs[0]
            rr = pp.add_run(text)
            rr.font.size = Pt(6.8 if i != 1 else 7.5)
            rr.bold = i < 2
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pp.paragraph_format.space_after = Pt(0)
            if j == 2:
                shade(cc, COL["paper"])
    widths(t, [w / 4 - 0.02] * 4)

    label("KEY STATISTICS")
    stats = [("Symbol", f"{short_exchange(c.get('exchange'))}: {ctx.ticker}"),
             ("52-Week Range", f"{ctx.p(c['week52_low'])} – {ctx.p(c['week52_high'])}"),
             ("YTD Performance", fmt.pct(c["ytd_return"])),
             ("Market Cap", fmt.money(c["market_cap"], ctx.cur)),
             ("Dividend Yield", fmt.pct(c["dividend_yield"])),
             ("NTM P/E", fmt.multiple(c["ntm_pe"])),
             ("NTM EV/EBITDA", fmt.multiple(c["ntm_ev_ebitda"])),
             ("ROE", fmt.pct(c["roe"])), ("ROA", fmt.pct(c["roa"])), ("ROIC", fmt.pct(c["roic"]))]
    t2 = cell.add_table(rows=len(stats), cols=2)
    table_borders(t2, horizontal=COL["grid"])
    cell_margins(t2, 10, 10, 20, 20)
    for i, (k, v) in enumerate(stats):
        for j, text in enumerate((k, v)):
            pp = t2.rows[i].cells[j].paragraphs[0]
            rr = pp.add_run(text)
            rr.font.size = Pt(7)
            rr.bold = j == 0
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            pp.paragraph_format.space_after = Pt(0)
    widths(t2, [w * 0.52, w * 0.46])

    et = m["eps_table"]
    years = list(et["years"])
    basis = "Adj." if et["basis"] == "adjusted" else "GAAP"
    fye = (ctx.company["edgar"].get("fiscal_year_end_mmdd") or "1231")
    month = date(2000, int(fye[:2]), 1).strftime("%b")
    label(f"EPS ({basis})")
    rows = [[f"FY ({month})"] + [y.replace("FY", "") + ("A" if i == 0 else "E") for i, y in enumerate(years)]]
    for q in range(4):
        rows.append([f"Q{q + 1}"] + [fmt_eps(et["years"][y]["quarters"][q], et["years"][y]["kinds"][q]) for y in years])
        rows.append(["  YoY"] + ([""] + [fmt.pct((et["years"][y].get("yoy_quarters") or [None] * 4)[q]) for y in years[1:]]))
    rows.append(["Year"] + [fmt.price(et["years"][y]["year"]).replace("$", "") if et["years"][y]["year"] is not None else "N/A"
                            for y in years])
    t3 = cell.add_table(rows=len(rows), cols=4)
    table_borders(t3, horizontal=COL["grid"], outer=COL["ink"])
    cell_margins(t3, 5, 5, 20, 20)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cc = t3.rows[i].cells[j]
            pp = cc.paragraphs[0]
            rr = pp.add_run(text)
            yoy = row[0].strip() == "YoY"
            rr.font.size = Pt(6 if yoy else 6.8)
            rr.bold = i == 0 or i == len(rows) - 1 or (j == 0 and not yoy)
            rr.italic = yoy
            rr.font.color.rgb = rgb(COL["text_secondary"] if yoy else COL["ink"])
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            pp.paragraph_format.space_after = Pt(0)
            if i == 0:
                shade(cc, COL["paper"])
    widths(t3, [w * 0.25, w * 0.24, w * 0.24, w * 0.24])

    label("ANALYST")
    for text, bold in ((AUTHOR["name"], True), (AUTHOR["email"], False)):
        pp = cell.add_paragraph()
        rr = pp.add_run(text)
        rr.font.size = Pt(7.5)
        rr.bold = bold
        pp.paragraph_format.space_after = Pt(0)


def fmt_eps(v, kind) -> str:
    if v is None:
        return "N/A"
    s = f"{v:,.2f}"
    return s + ("*" if kind and "override" in kind else "")


def short_exchange(x: str | None) -> str:
    if not x:
        return ""
    x = x.upper()
    for k in ("NASDAQ", "NYSE"):
        if k in x:
            return k
    return x


# ---------------------------------------------------------------- toc + disclaimer

def toc(doc, headings: list[tuple[int, str]]) -> None:
    t = plain_heading(doc, "Table of Contents", size=12.5)
    para_border(t, "bottom", COL["grid"], 4)
    p = doc.add_paragraph()
    placeholder = "\n".join(("    " if lvl == 2 else "") + h for lvl, h in headings)
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', placeholder or "Right-click to update the table of contents.")


def disclaimer(doc, holdings: str) -> None:
    text = (GUIDES_DIR / "disclaimer.md").read_text()
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    text = re.sub(r"^# .*\n", "", text)
    text = text.replace("{{HOLDINGS_DISCLOSURE}}", holdings or "[VERIFY: holdings disclosure]")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for block in re.split(r"\n\s*\n", text.strip()):
        lines = block.strip().splitlines()
        if lines and lines[0].startswith("## "):
            plain_heading(doc, lines[0][3:].strip())
            lines = lines[1:]
        if lines:
            p = body_para(doc, " ".join(x.strip() for x in lines), size=8)
            p.paragraph_format.space_after = Pt(5)


def plain_heading(doc, text: str, size: float = 11):
    """Heading look without a Heading style (kept out of the TOC)."""
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    r.font.name, r.font.size, r.bold = HEAD_FONT, Pt(size), True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    keep_with_next(p)
    return p


# ---------------------------------------------------------------- build

def section_files(ctx: Ctx) -> list[Path]:
    return sorted((ctx.dir / "sections").glob("*.md"))


def build(ticker: str, log=print, update_on_open: bool = True) -> Path:
    ctx = Ctx(ticker)
    files = section_files(ctx)
    if not files:
        raise FileNotFoundError(f"No sections in {ctx.dir / 'sections'}. Run: uv run erb scaffold {ticker}")
    cover_file = next((f for f in files if f.name.startswith("00")), None)
    front, cover_body = split_front(cover_file.read_text()) if cover_file else ({}, "")
    ctx.front = front
    report_date = fmt_date(front.get("report_date") or date.today().isoformat())

    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(MARGIN))
    sec.top_margin, sec.bottom_margin = Inches(0.55), Inches(0.6)
    sec.header_distance, sec.footer_distance = Inches(0.3), Inches(0.3)
    header_footer(sec, ctx, report_date)

    cover(doc, ctx, front, cover_body, report_date)

    bodies = [(f, split_front(f.read_text())[1]) for f in files if f is not cover_file]
    headings = []
    for _, text in bodies:
        clean = re.sub(r"<!--.*?-->", "", text, flags=re.S)
        for ln in clean.splitlines():
            if ln.startswith("# "):
                headings.append((1, SOURCE_TAG.sub("", ln[2:].strip()).upper()))
            elif ln.startswith("## "):
                headings.append((2, SOURCE_TAG.sub("", ln[3:].strip())))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    toc(doc, headings)

    for f, text in bodies:
        render_md(doc, text, ctx, scenario_headings=f.name.startswith("01"))

    disclaimer(doc, front.get("holdings_disclosure", ""))

    # Ask Word to refresh the TOC/page fields on open (erb build --word does it headlessly instead)
    if update_on_open:
        settings = doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)

    name = re.sub(r"[^A-Za-z0-9]+", "-", ctx.name).strip("-")
    out = ctx.dir / f"CI_{name}_{ctx.ticker}_Initiating-Coverage_{date.fromisoformat(str(front.get('report_date') or date.today()))}.docx"
    doc.save(out)
    for t in sorted(set(ctx.missing_tokens)):
        log(f"NOTE: exhibit not available: {t}")
    log(f"Wrote {out}")
    return out
