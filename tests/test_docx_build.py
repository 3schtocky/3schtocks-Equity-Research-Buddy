"""Build a report from the META fixtures if present (skips on a fresh clone)."""
import re

import pytest
from docx import Document

from erb import docx_build
from erb.config import coverage_dir


def test_inline_markup_strips_source_tags_and_flags_verify():
    doc = Document()
    p = doc.add_paragraph()
    docx_build.add_inline(p, "Revenue rose 22.2% [S3] to **$201.0 bn** [S3, S4]. [VERIFY: CEO name]")
    assert p.text == "Revenue rose 22.2% to $201.0 bn. [VERIFY: CEO name]"
    assert any(r.bold for r in p.runs if r.text == "$201.0 bn")
    assert any(r.font.highlight_color is not None for r in p.runs if r.text.startswith("[VERIFY"))


def test_front_matter_split():
    fm, body = docx_build.split_front("---\ntagline: Hi!\n---\n- bullet\n")
    assert fm == {"tagline": "Hi!"} and body.strip() == "- bullet"


@pytest.mark.skipif(not (coverage_dir("META") / "model.json").exists()
                    or not (coverage_dir("META") / "sections").exists(), reason="needs META coverage fixtures")
def test_full_build_has_all_sections_and_no_unrendered_tokens(tmp_path):
    out = docx_build.build("META", log=lambda *_: None)
    d = Document(out)
    text = "\n".join(p.text for p in d.paragraphs)
    for h in ("INVESTMENT SUMMARY", "VALUATION ANALYSIS", "FINANCIAL ANALYSIS", "DISCLAIMER"):
        assert h in text.upper()
    assert not re.search(r"\{\{\s*\w+\s*\}\}", text)
    assert not re.search(r"\[S\d+\]", text)
    assert len(d.tables) >= 6
